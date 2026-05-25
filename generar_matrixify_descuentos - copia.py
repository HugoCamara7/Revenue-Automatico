from __future__ import annotations

import argparse
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


MATRIXIFY_REQUIRED_HEADERS = [
    "ID",
    "Handle",
    "Command",
    "Variant Inventory Item ID",
    "Variant ID",
    "Variant Command",
    "Variant SKU",
    "Variant Price",
    "Variant Compare At Price",
]


@dataclass
class DiscountColumn:
    index: int
    label: str
    header: str
    kind: str


def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def as_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "").replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    return number / 100 if "%" in str(value) or number > 1 else number


def money(value: Any) -> float | None:
    number = as_number(value)
    return round(number, 2) if number is not None else None


def clean_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "-", name).strip()
    cleaned = re.sub(r"\s+", " ", cleaned) or "Descuento"
    cleaned = cleaned[:31].strip(" -")
    base = cleaned
    suffix = 2
    while cleaned in used:
        tail = f" {suffix}"
        cleaned = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    used.add(cleaned)
    return cleaned


def find_header_row(ws: Worksheet, required: list[str], scan_rows: int = 30) -> tuple[int, dict[str, int]]:
    required_norm = {normalize(header): header for header in required}
    for row in range(1, min(ws.max_row, scan_rows) + 1):
        found: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            cell_norm = normalize(ws.cell(row=row, column=col).value)
            if cell_norm in required_norm:
                found[required_norm[cell_norm]] = col
        if all(header in found for header in required):
            return row, found
    raise ValueError(f"No encontre encabezados requeridos: {', '.join(required)}")


def find_revenue_layout(ws: Worksheet) -> tuple[int, dict[str, int], list[DiscountColumn]]:
    header_row, columns = find_header_row(ws, ["ID PRODUCTO"], scan_rows=30)
    for col in range(1, ws.max_column + 1):
        header = normalize(ws.cell(row=header_row, column=col).value)
        if header:
            columns[header] = col

    discount_columns: list[DiscountColumn] = []
    for col in range(1, ws.max_column + 1):
        header_value = ws.cell(row=header_row, column=col).value
        header = normalize(header_value)
        if "DCTO" not in header and "DESCUENTO" not in header:
            continue

        top_labels = []
        for row in range(1, header_row):
            value = ws.cell(row=row, column=col).value
            if value not in (None, ""):
                top_labels.append(str(value).strip())

        if not top_labels:
            left = col - 1
            while left >= 1 and not top_labels:
                for row in range(1, header_row):
                    value = ws.cell(row=row, column=left).value
                    if value not in (None, ""):
                        top_labels.append(str(value).strip())
                left -= 1

        label = " - ".join(top_labels + [str(header_value).strip()])
        kind = "resto_mes" if "ANT" in header else "programado"
        discount_columns.append(DiscountColumn(col, label, str(header_value).strip(), kind))

    if not discount_columns:
        raise ValueError("No encontre columnas de descuento en el archivo Revenue.")
    return header_row, columns, discount_columns


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def build_discount_workbook(matrixify_path: Path, revenue_path: Path, output_path: Path) -> Path:
    matrix_wb = load_workbook(matrixify_path)
    matrix_ws = matrix_wb["Products"] if "Products" in matrix_wb.sheetnames else matrix_wb.active
    matrix_header_row, matrix_cols = find_header_row(matrix_ws, MATRIXIFY_REQUIRED_HEADERS)

    revenue_wb = load_workbook(revenue_path, data_only=True)
    revenue_ws = revenue_wb.active
    revenue_header_row, revenue_cols, discount_columns = find_revenue_layout(revenue_ws)

    sku_col = matrix_cols["Variant SKU"]
    group_col = matrix_cols.get("ID") or matrix_cols.get("Handle") or sku_col
    group_to_rows: dict[str, list[int]] = {}
    sku_to_group: dict[str, str] = {}

    for row in range(matrix_header_row + 1, matrix_ws.max_row + 1):
        sku = matrix_ws.cell(row=row, column=sku_col).value
        group_key = str(matrix_ws.cell(row=row, column=group_col).value or sku or "").strip()
        if not group_key:
            continue
        group_to_rows.setdefault(group_key, []).append(row)
        if sku not in (None, ""):
            sku_to_group[str(sku).strip()] = group_key

    discounts_by_col: dict[int, dict[str, float]] = {col.index: {} for col in discount_columns}
    missing_rows: list[list[Any]] = [["Hoja", "Variant SKU / ID PRODUCTO", "Descuento"]]
    invalid_rows: list[list[Any]] = [["Variant SKU / ID PRODUCTO", "Columna Revenue", "Valor leido"]]
    id_col = revenue_cols["ID PRODUCTO"]

    for row in range(revenue_header_row + 1, revenue_ws.max_row + 1):
        product_id = revenue_ws.cell(row=row, column=id_col).value
        if product_id in (None, ""):
            continue
        sku = str(product_id).strip()
        for discount_col in discount_columns:
            discount = as_number(revenue_ws.cell(row=row, column=discount_col.index).value)
            if discount is None:
                continue
            if discount < 0 or discount >= 1:
                invalid_rows.append([sku, discount_col.label, discount])
                continue
            discounts_by_col[discount_col.index][sku] = discount

    output_wb = Workbook()
    output_wb.remove(output_wb.active)
    used_names: set[str] = set()
    total_rows = matrix_ws.max_row - matrix_header_row
    summary_rows = [
        [
            "Hoja",
            "Tipo",
            "Columna Revenue",
            "SKUs con descuento",
            "Productos/Modelo-color con descuento",
            "Filas Matrixify generadas",
            "Filas sin descuento",
            "SKUs Revenue no encontrados",
        ]
    ]

    for discount_col in discount_columns:
        discounts = discounts_by_col[discount_col.index]
        if not discounts:
            continue

        sheet_name = clean_sheet_name(discount_col.label, used_names)
        out_ws = output_wb.create_sheet(sheet_name)

        for col in range(1, len(MATRIXIFY_REQUIRED_HEADERS) + 1):
            source = matrix_ws.cell(row=matrix_header_row, column=col)
            target = out_ws.cell(row=1, column=col, value=source.value)
            copy_cell_style(source, target)
            out_ws.column_dimensions[target.column_letter].width = max(
                matrix_ws.column_dimensions[source.column_letter].width or 12,
                12,
            )

        groups_to_discount: dict[str, float] = {}
        missing = 0
        for sku, discount in discounts.items():
            group_key = sku_to_group.get(sku)
            if not group_key:
                missing += 1
                missing_rows.append([sheet_name, sku, discount])
                continue
            groups_to_discount[group_key] = max(groups_to_discount.get(group_key, 0), discount)

        out_row = 2
        discounted_output_rows = 0
        for matrix_row in range(matrix_header_row + 1, matrix_ws.max_row + 1):
            group_key = str(matrix_ws.cell(row=matrix_row, column=group_col).value or "").strip()
            discount = groups_to_discount.get(group_key)
            original_compare = money(matrix_ws.cell(row=matrix_row, column=matrix_cols["Variant Compare At Price"]).value)
            current_price = money(matrix_ws.cell(row=matrix_row, column=matrix_cols["Variant Price"]).value)
            original_price = original_compare or current_price

            if original_price is None:
                new_price = matrix_ws.cell(row=matrix_row, column=matrix_cols["Variant Price"]).value
                compare_at = None
            elif discount is None or discount <= 0:
                new_price = original_price
                compare_at = None
            else:
                new_price = round(original_price * (1 - discount), 2)
                compare_at = original_price if new_price != original_price else None
                if compare_at is not None:
                    discounted_output_rows += 1

            for col in range(1, len(MATRIXIFY_REQUIRED_HEADERS) + 1):
                source = matrix_ws.cell(row=matrix_row, column=col)
                value = source.value
                if col == matrix_cols["Variant Price"]:
                    value = new_price
                elif col == matrix_cols["Variant Compare At Price"]:
                    value = compare_at
                target = out_ws.cell(row=out_row, column=col, value=value)
                copy_cell_style(source, target)
                if col in (matrix_cols["Variant Price"], matrix_cols["Variant Compare At Price"]):
                    target.number_format = "0.00"
            out_row += 1

        out_ws.freeze_panes = "A2"
        out_ws.auto_filter.ref = out_ws.dimensions
        summary_rows.append(
            [
                sheet_name,
                "Resto del mes" if discount_col.kind == "resto_mes" else "Programado",
                discount_col.label,
                len(discounts),
                len(groups_to_discount),
                total_rows,
                total_rows - discounted_output_rows,
                missing,
            ]
        )

    summary_ws = output_wb.create_sheet("Resumen", 0)
    for row in summary_rows:
        summary_ws.append(row)
    summary_ws.freeze_panes = "A2"
    summary_ws.auto_filter.ref = summary_ws.dimensions
    for col in range(1, 9):
        summary_ws.column_dimensions[summary_ws.cell(row=1, column=col).column_letter].width = 28

    if len(missing_rows) > 1:
        missing_ws = output_wb.create_sheet("No encontrados")
        for row in missing_rows:
            missing_ws.append(row)
        missing_ws.freeze_panes = "A2"
        missing_ws.auto_filter.ref = missing_ws.dimensions

    if len(invalid_rows) > 1:
        invalid_ws = output_wb.create_sheet("Descuentos invalidos")
        for row in invalid_rows:
            invalid_ws.append(row)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera Matrixify de descuentos por hoja/campana.")
    parser.add_argument("--matrixify", type=Path, required=True)
    parser.add_argument("--revenue", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(build_discount_workbook(args.matrixify, args.revenue, args.output))


if __name__ == "__main__":
    main()
