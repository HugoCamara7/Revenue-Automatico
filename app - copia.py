from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import streamlit as st

from generar_matrixify_descuentos import build_discount_workbook


SITES = {
    "Columbia.pe": {"brand": "COLUMBIA", "vendor": "columbiape", "output": "matrixify_columbia_descuentos.xlsx"},
    "Hush Puppies.pe": {"brand": "HUSH PUPPIES", "vendor": "hushpuppiespe", "output": "matrixify_hushpuppies_descuentos.xlsx"},
    "Rockford.pe": {"brand": "ROCKFORD", "vendor": "rockfordpe", "output": "matrixify_rockford_descuentos.xlsx"},
    "Supermall.pe": {"brand": "MULTIMARCA", "vendor": "supermallpe", "output": "matrixify_supermall_descuentos.xlsx"},
}


st.set_page_config(page_title="Matrixify Descuentos", layout="wide")

st.markdown(
    """
    <style>
    .stApp { background: #ffffff; color: #001f4f; }
    [data-testid="stSidebar"] { background: #eef2f7; }
    .main-block { max-width: 960px; margin: 0 auto; padding-top: 28px; }
    .brand-logo { font-size: 58px; line-height: 1; font-weight: 800; color: #15329b; letter-spacing: 1px; }
    .brand-subtitle { color: #15329b; font-size: 15px; font-weight: 700; letter-spacing: 7px; margin-bottom: 64px; }
    .shopify-badge {
        height: 84px; width: 84px; border-radius: 12px; background: #95bf47; color: #ffffff;
        display: flex; align-items: center; justify-content: center; font-size: 48px; font-weight: 800;
        transform: rotate(4deg); margin-left: auto;
    }
    .info-box { background: #eef6ff; border: 1px solid #c3ddff; border-radius: 8px; padding: 20px 24px; margin: 34px 0 40px 0; }
    .section-card { border: 1px solid #d7e1ef; border-radius: 8px; padding: 28px 24px; margin-bottom: 18px; }
    .upload-card { border: 1px dashed #8ab4ff; border-radius: 8px; padding: 20px 18px; margin-bottom: 16px; }
    </style>
    """,
    unsafe_allow_html=True,
)


def save_upload(uploaded_file, folder: Path) -> Path:
    path = folder / uploaded_file.name
    path.write_bytes(uploaded_file.getbuffer())
    return path


def read_summary(output_path: Path) -> tuple[list[str], list[list[object]], int]:
    workbook = openpyxl.load_workbook(output_path, data_only=True, read_only=True)
    summary = workbook["Resumen"]
    headers = [summary.cell(row=1, column=col).value for col in range(1, summary.max_column + 1)]
    rows = [
        [summary.cell(row=row, column=col).value for col in range(1, summary.max_column + 1)]
        for row in range(2, summary.max_row + 1)
    ]
    missing_count = workbook["No encontrados"].max_row - 1 if "No encontrados" in workbook.sheetnames else 0
    workbook.close()
    return headers, rows, missing_count


with st.sidebar:
    site_name = st.selectbox("Sitio destino", list(SITES.keys()))
    site = SITES[site_name]
    st.markdown("**Marcas permitidas**")
    st.write(site["brand"])
    st.caption(f"Vendor: {site['vendor']} | Salida: {site['output']}")

st.markdown('<div class="main-block">', unsafe_allow_html=True)
top_left, top_right = st.columns([3, 1])
with top_left:
    st.markdown('<div class="brand-logo">FORUS</div>', unsafe_allow_html=True)
    st.markdown('<div class="brand-subtitle">CONSUMER FANATIC</div>', unsafe_allow_html=True)
with top_right:
    st.markdown('<div class="shopify-badge">S</div>', unsafe_allow_html=True)

st.header(f"Matrixify {site_name} - Shopify")
st.write(
    "Sube el Revenue de descuentos y el ultimo catalogo Matrixify del sitio. "
    "La app genera hojas completas por campana y mantiene limpio el formato de precios."
)
st.markdown(
    """
    <div class="info-box">
      <strong>Flujo obligatorio:</strong><br>
      Elige el sitio destino, sube el Revenue comercial y sube el ultimo catalogo Matrixify del mismo sitio.
      Cada hoja generada contiene todo el catalogo: los productos sin descuento mantienen precio original y Compare At vacio.
    </div>
    """,
    unsafe_allow_html=True,
)

st.markdown('<div class="section-card"><h3>Cargar archivos obligatorios</h3></div>', unsafe_allow_html=True)
st.markdown('<div class="upload-card">', unsafe_allow_html=True)
revenue_file = st.file_uploader("1. Subir input comercial / Revenue", type=["xlsx", "xls"])
st.markdown("</div>", unsafe_allow_html=True)
st.markdown('<div class="upload-card">', unsafe_allow_html=True)
matrixify_file = st.file_uploader(f"2. Subir ultimo catalogo Matrixify de {site_name}", type=["xlsx", "xls"])
st.markdown("</div>", unsafe_allow_html=True)

with st.expander("Regla de calculo"):
    st.write(
        "El match se hace con `Revenue.ID PRODUCTO` contra `Matrixify.Variant SKU`. "
        "Si un SKU tiene descuento, la app aplica el descuento a todas las variantes del mismo `ID` de producto."
    )
    st.write(
        "`Variant Price` queda con formato uniforme de 2 decimales. "
        "`Variant Compare At Price` solo se llena si hay descuento real; si no, queda vacio."
    )

if st.button("Generar Matrixify de descuentos", type="primary", disabled=not matrixify_file or not revenue_file, use_container_width=True):
    with st.spinner("Procesando archivos..."):
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                workdir = Path(temp_dir)
                revenue_path = save_upload(revenue_file, workdir)
                matrixify_path = save_upload(matrixify_file, workdir)
                output_path = workdir / site["output"]
                build_discount_workbook(matrixify_path, revenue_path, output_path)
                headers, rows, missing_count = read_summary(output_path)
                output_bytes = output_path.read_bytes()

            st.success("Archivo generado correctamente.")
            st.subheader("Resumen")
            st.dataframe([dict(zip(headers, row)) for row in rows], hide_index=True, use_container_width=True)
            if missing_count:
                st.warning(f"{missing_count} SKUs del Revenue no se encontraron; el catalogo igual se genero completo.")
            st.download_button(
                "Descargar Matrixify generado",
                data=output_bytes,
                file_name=site["output"],
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        except Exception as exc:
            st.error(f"No pude generar el archivo: {exc}")

st.markdown("</div>", unsafe_allow_html=True)
